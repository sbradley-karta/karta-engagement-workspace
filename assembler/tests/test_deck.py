import json
import os
from pathlib import Path

import pytest

from karta_assembler.deck import (
    ApprovedValuesError,
    EMPTY_RISKS_TEXT,
    STATUS_THEME_COLOR,
    build_deck,
    deck_filename,
    fit_to_budget,
    long_date,
    short_date,
    validate_approved_values,
    wrapped_lines,
)
from karta_assembler.status import status_period


def approved(**over):
    p = status_period("2026-09-04")
    av = {
        "schema_version": "1.0",
        "engagement": {"engagement_id": "eng-test", "client": "TestCo", "project_name": "Planning Pilot"},
        "period": {"start": p.start.isoformat(), "end": p.end.isoformat(), "label": p.label},
        "as_of": "2026-09-04",
        "project_managers": ["A. Reviewer", "B. Lead"],
        "milestones": [
            {"name": "Kickoff", "completed": True, "completed_at": "2026-06-01T12:00:00Z", "original_date": "2026-06-01", "new_date": None, "status": "Complete", "variance_days": None, "source": {"system": "asana", "id": "1"}},
            {"name": "Sprint review", "completed": False, "completed_at": None, "original_date": "2026-08-14", "new_date": "2026-09-09", "status": "At Risk", "variance_days": 26, "source": {"system": "asana", "id": "2"}},
            {"name": "UAT start", "completed": False, "completed_at": None, "original_date": "2026-09-21", "new_date": None, "status": "On Track", "variance_days": None, "source": {"system": "asana", "id": "3"}},
        ],
        "assessments": {
            "scope_schedule": {"status": "Yellow", "rationale": "One review re-planned."},
            "resources": {"status": "Green", "rationale": "Staffed."},
            "data": {"status": "Green", "rationale": "Loads complete."},
        },
        "overall": {"status": "Yellow", "derivation": "worst_of"},
        "accomplishments": ["Completed data hub load", "Reviewed design with the client"],
        "planned_activities": ["Sprint review on September 9", "Begin UAT script drafting"],
        "raid": [{"type": "Risk", "priority": "High", "description": "Source data refresh timing", "mitigation": "Confirm schedule with IT", "status": "In Progress"}],
        "source_gaps": [],
        "sources": [{"system": "asana", "retrieved_at": "2026-09-04T15:00:00Z"}],
        "approval": {"reviewer_name": "A. Reviewer", "reviewer_email": "reviewer@example.com", "approved_at": "2026-09-04T16:00:00Z", "version": 1},
    }
    av.update(over)
    return av


def test_dates_and_filename():
    assert short_date("2026-09-04") == "9/4"
    assert short_date(None) == ""
    assert long_date("2026-09-04") == "September 4, 2026"
    assert deck_filename("DoTERRA", "2026-08-26") == "KCG_Status DoTERRA 8-26-2026.pptx"


def test_line_budget():
    assert wrapped_lines("x" * 48) == 1 and wrapped_lines("x" * 49) == 2
    items = ["a" * 48] * 13 + ["b" * 100] + ["c"]
    kept, dropped = fit_to_budget(items)
    assert len(kept) == 13 and dropped == ["b" * 100, "c"]


def test_theme_colors_cover_every_status():
    assert set(STATUS_THEME_COLOR) == {"Green", "Yellow", "Red", "Unknown", "Not assessed"}


def test_validate_accepts_consistent_document():
    validate_approved_values(approved())


def test_validate_refuses_draft_without_approval():
    av = approved(); av["approval"] = {}
    with pytest.raises(ApprovedValuesError, match="named approval"):
        validate_approved_values(av)


def test_validate_refuses_wrong_milestone_status():
    av = approved(); av["milestones"][1]["status"] = "On Track"
    with pytest.raises(ApprovedValuesError, match="derives 'At Risk'"):
        validate_approved_values(av)


def test_validate_refuses_wrong_rollup():
    av = approved(); av["overall"]["status"] = "Green"
    with pytest.raises(ApprovedValuesError, match="worst-of"):
        validate_approved_values(av)


def test_validate_refuses_green_with_not_assessed():
    av = approved()
    av["assessments"]["data"]["status"] = "Not assessed"; av["overall"]["status"] = "Yellow"
    validate_approved_values(av)  # Yellow beats Unknown, consistent
    av["assessments"]["scope_schedule"]["status"] = "Green"; av["overall"]["status"] = "Green"
    with pytest.raises(ApprovedValuesError):
        validate_approved_values(av)


def test_validate_refuses_oversized():
    av = approved(); av["accomplishments"] = ["x" * 600_000]
    with pytest.raises(ApprovedValuesError, match="limit"):
        validate_approved_values(av)


BASE = os.environ.get("KCG_BASE_DECK")


@pytest.mark.skipif(not BASE, reason="set KCG_BASE_DECK to a real template or deck to run the build")
def test_build_against_real_base_deck(tmp_path):
    from pptx import Presentation
    data, manifest = build_deck(Path(BASE).read_bytes(), approved())
    assert manifest.filename == "KCG_Status TestCo 9-4-2026.pptx"
    prs = Presentation(str(_write(tmp_path, data)))
    s = [x for x in prs.slides if any(sh.name == "Table 14" for sh in x.shapes)][0]
    by = {sh.name: sh for sh in s.shapes}
    assert by["Table 9"].table.cell(0, 0).text == "Overall Status: Yellow"
    assert by["Table 14"].table.cell(2, 0).text == "Sprint review" and by["Table 14"].table.cell(2, 2).text == "9/9"
    assert by["Table 14"].table.cell(4, 0).text == ""
    assert by["Table 19"].table.cell(1, 0).text == "Risk (High)"
    paras = by["Table 16"].table.cell(1, 0).text_frame.paragraphs
    assert [p.text for p in paras] == ["Completed data hub load", "Reviewed design with the client"]


@pytest.mark.skipif(not BASE, reason="set KCG_BASE_DECK")
def test_empty_raid_phrasing(tmp_path):
    from pptx import Presentation
    data, _ = build_deck(Path(BASE).read_bytes(), approved(raid=[]))
    prs = Presentation(str(_write(tmp_path, data)))
    s = [x for x in prs.slides if any(sh.name == "Table 19" for sh in x.shapes)][0]
    t19 = [sh for sh in s.shapes if sh.name == "Table 19"][0].table
    assert t19.cell(1, 1).text == EMPTY_RISKS_TEXT and t19.cell(2, 1).text == ""


def _write(tmp_path, data):
    p = tmp_path / "out.pptx"; p.write_bytes(data); return p
